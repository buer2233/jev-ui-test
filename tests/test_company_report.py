"""公司规范格式测试报告：**离线**用合成结果打一遍。

为什么必须离线可测：这份报告是交付物，而它唯一的输入是 Allure 结果目录——
真跑一次 NL 用例要花钱、要 E9 环境、还要四十多分钟。这里用**手写的
`allure-results` 夹具**（`_report_fixtures.py`）把每种边界打一遍：

  · 通过 / 失败 / 跳过三条路径的措辞与配图；
  · **重跑**（同一 historyId 多份结果）只占一行、并把"重跑 N 次"写进备注；
  · 跳过用例**没有 `nl_case` 参数**（它没进过 runner）时仍能认领到用例 id；
  · 断言截图缺失时写「（无截图）」而**不是崩掉**；
  · 版式对照校验**真的会红**（用一份被改坏的"模版"验）。

夹具的 `cases/` 是**临时目录**里的最小 YAML，不是仓库里的 `cases/`——
单测不该因为有人改了真实用例而变红或变绿。
"""

from datetime import datetime

import openpyxl
import pytest
from _report_fixtures import write_results

from jev_ultrafast.framework import company_report as cr
from jev_ultrafast.framework.company_report import (
    NO_SCREENSHOT,
    ReportCase,
    build_model,
    compare_with_template,
    load_results,
    write_report,
)


@pytest.fixture
def fixture(tmp_path):
    return write_results(tmp_path)


@pytest.fixture
def model(fixture):
    return build_model(fixture.results, fixture.cases)


def _by_name(model, name):
    return next(case for case in model.cases if case.name == name)


# ---------------------------------------------------------------- 读取层


def test_retries_collapse_into_one_row_and_are_reported(fixture, model):
    """同一 historyId 的多份结果只占**一行**，保留最后一次，次数记进 retries。

    重跑才过的用例必须能被看出来（AGENTS.md：把重跑兜住的报成"通过"是不准确的）。
    """
    results = load_results(fixture.results)
    assert len(results) == 5, "重跑的两份结果必须合成一行"

    retried = _by_name(model, "重跑才过的用例")
    assert retried.status == "passed"
    assert retried.retries == 1
    assert "重跑 1 次" in retried.notes


def test_results_are_ordered_by_start_time(model):
    """顺序按**开始时间**，不是文件名字典序——后者会把报告里的顺序打乱成看不懂的样子。"""
    assert [case.name for case in model.cases] == [
        "通过的用例", "失败的用例", "跳过用例", "重跑才过的用例", "没有截图的用例",
    ]


def test_case_id_falls_back_to_the_parameterized_name(fixture, model):
    """跳过的用例没有 `nl_case` 参数（没进过 runner）→ 按名字尾部的 id 认领。

    认领不到的话，它在报告里会变成"没有类别、没有 check项"的孤儿——
    而那正是**最需要写清理由**的那一条用例。
    """
    skipped = _by_name(model, "跳过用例")
    assert skipped.case_id == "e9-skip-case"
    assert skipped.check == "跳过用例的 check 项"
    assert skipped.category == "负向·环境前提不成立"
    assert "环境前提不成立" in skipped.actual


# ---------------------------------------------------------------- 组装层


def test_design_metadata_comes_from_the_yaml_report_block(model):
    """测试项类别 / check项 / 预期 取自用例 YAML 的 `report:` 块（按 id 对齐）。"""
    case = _by_name(model, "通过的用例")
    assert case.category == "发起·提交"
    assert case.check == "提交后表单关闭并回到列表"
    assert case.expect == "点「提 交」后表单表单关闭、回到流程名称列表"
    assert case.outline == "提交后表单关闭并回到列表（发起·提交）"   # 模版的写法


def test_actual_text_is_built_from_assertion_details_not_from_a_model(model):
    """实际结果逐条来自「断言汇总」的 detail——**不经过任何模型**，改不了也编不出来。"""
    passed = _by_name(model, "通过的用例")
    assert passed.actual.startswith("【通过】")
    assert "URL 期望包含 'requestid='" in passed.actual
    assert "浏览器动作 7 个" in passed.actual        # 执行摘要里的计数也带上


def test_failed_case_keeps_the_failure_message(model):
    failed = _by_name(model, "失败的用例")
    assert failed.actual.startswith("【失败】")
    assert "页面文本期望出现" in failed.actual
    assert "实际 0 次" in failed.actual


def test_missing_screenshot_is_filled_with_a_placeholder_not_a_crash(fixture, model, tmp_path):
    """没有断言截图 → `（无截图）`，**不是**空白、也不是崩掉。

    模版自己就有「（无截图）」这种行：它是"这一步没拍到"，不是"报告坏了"。
    """
    assert _by_name(model, "没有截图的用例").screenshot is None
    assert write_report(model, tmp_path / "报告.xlsx").is_file()


def test_environment_block_never_invents_values(model):
    """环境块只写**取到的东西**，取不到的一律「本次未采集」——不编。

    夹具里既没有 `environment.properties`，也没配 E9 地址：
    服务器那一栏就得如实写"未采集"，而客户端那一栏里生成时刻就能拿到的操作系统照常写。
    """
    server, client = cr._environment_rows(model)
    assert any("本次未采集" in line for line in server)
    assert any("Windows" in line or "Linux" in line or "Darwin" in line for line in client)


def test_environment_block_uses_the_runtime_case_not_the_first_loaded_one(fixture):
    """地址/模块取自**这批结果真正覆盖到的用例**，不是"加载到的第一条"。

    ⚠️ 这条防的是一个看不见的错：`cases/` 里同时躺着演示用例（没有 E9 地址），
    按"第一条"算出来的"被测系统"会是演示站——报告上完全看不出来是错的。
    """
    model = build_model(fixture.results, fixture.cases)
    assert model.environment["测试地址"] == "http://e9.example:8080"
    assert model.environment["研发模块"] == "流程提交"


# ---------------------------------------------------------------- 写出层


def test_write_report_produces_the_three_sheets_in_template_order(model, tmp_path):
    path = write_report(model, tmp_path / "报告.xlsx")
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == [cr.SHEET_PLAN, cr.SHEET_LIST, cr.SHEET_SHOTS]
    assert workbook[cr.SHEET_SHOTS].freeze_panes == "A2"      # 模版冻结表头


def test_screenshot_sheet_embeds_an_image_per_case_that_has_one(model, tmp_path):
    path = write_report(model, tmp_path / "报告.xlsx")
    shots = openpyxl.load_workbook(path)[cr.SHEET_SHOTS]
    # 夹具里 3 条有图（通过 / 失败 / 重跑），其余写「（无截图）」
    assert len(shots._images) == 3
    assert any(cell.value == NO_SCREENSHOT for row in shots.iter_rows() for cell in row)


def test_confirm_columns_skip_the_cases_that_never_ran(fixture, tmp_path):
    """确认列（√）只给**真跑过**的用例——给跳过的也打勾，等于把没做的活报成做了。"""
    model = build_model(fixture.results, fixture.cases)
    sheet = openpyxl.load_workbook(write_report(model, tmp_path / "报告.xlsx"))[cr.SHEET_LIST]

    rows = {sheet[f"C{row}"].value: row for row in range(10, 10 + len(model.cases))}
    assert sheet[f"E{rows['跳过用例的 check 项']}"].value in (None, "")
    assert sheet[f"E{rows['提交后表单关闭并回到列表']}"].value == "√"


def test_same_category_rows_share_one_merged_cell(fixture, tmp_path):
    """相邻同类别纵向合并——模版就是这么排的（不合并则每行重复一遍，读起来是噪声）。"""
    model = build_model(fixture.results, fixture.cases)
    sheet = openpyxl.load_workbook(write_report(model, tmp_path / "报告.xlsx"))[cr.SHEET_LIST]
    merged = {str(rng) for rng in sheet.merged_cells.ranges}
    # 前两条用例同属「发起·提交」→ B10:B11 应合并
    assert "B10:B11" in merged


# ---------------------------------------------------------------- 对照层


def test_compare_with_template_flags_a_structural_difference(fixture, tmp_path):
    """对照校验必须**真的会红**：拿一份版式被改坏的"模版"去比，它得报出来。

    没有这条，校验只证明"比过了"，不证明"比得对"。
    """
    model = build_model(fixture.results, fixture.cases)
    ours = write_report(model, tmp_path / "报告.xlsx")
    reference = tmp_path / "参照.xlsx"
    openpyxl.load_workbook(ours).save(reference)

    assert all(check.ok for check in compare_with_template(ours, reference))

    workbook = openpyxl.load_workbook(reference)
    workbook[cr.SHEET_LIST]["C7"] = "被改坏的列名"
    workbook.save(reference)

    checks = compare_with_template(ours, reference)
    assert any(not check.ok for check in checks)
    assert "[不一致]" in cr.format_checks(checks)
    assert "[通过]" in cr.format_checks(checks)


def test_default_output_path_never_overwrites(tmp_path):
    """一次一份、不覆盖——与 Allure 报告同一个理由（覆盖了就没法对比上一次）。"""
    stamp = datetime(2026, 10, 10, 12, 0, 0)
    first = cr.default_output_path(tmp_path, stamp)
    first.write_text("x", encoding="utf-8")
    second = cr.default_output_path(tmp_path, stamp)
    assert first != second
    assert second.name.startswith(first.stem)


def test_report_case_code_and_verdict_use_the_template_wording():
    """用例编号是 `T01` 这种形状、结论是【通过】这种措辞（与模版一致）。"""
    case = ReportCase(index=1, case_id="x", name="n", category="c", check="k", expect="e",
                      status="passed", actual="a", notes="n")
    assert case.code == "T01"
    assert case.verdict == "【通过】"


def test_every_allure_status_has_a_chinese_label():
    """Allure 的四种状态都要有中文标签——漏一个就会在报告里显示成【未知】。"""
    for status in ("passed", "failed", "broken", "skipped"):
        assert status in cr.STATUS_LABELS


def test_long_cells_are_truncated_with_a_visible_marker():
    """截断必须留痕：不然读报告的人以为原文就这么短。"""
    text = cr._short("甲" * (cr.CELL_TEXT_LIMIT + 50))
    assert text.endswith("全文见 Allure 报告）")
    assert len(text) < cr.CELL_TEXT_LIMIT + 40

def test_skipped_cases_are_not_reported_as_missing_screenshots(fixture, tmp_path):
    """跳过用例没有图**不算问题**（它没跑过，模版自己也有「（无截图）」的行）。

    把它也报出来的代价是假警报：每轮都响一次，真缺图那一次就没人看了。
    """
    summary = cr.export(fixture.results, tmp_path / "报告.xlsx", cases_dir=fixture.cases)
    assert "T03" not in summary.without_screenshot        # T03 是跳过那条
    assert summary.without_screenshot == ["T05"]          # 只有"跑过但没图"的那条


# ------------------------------------------------- 随仓库分发的公司模版

def test_bundled_template_ships_with_the_repo():
    """模版必须**随仓库分发**：新克隆的工作区要有参照物，否则版式校验无从谈起。

    它含真实项目信息与内网截图，是**已确认的取舍**（AGENTS.md「凭据与仓库边界」的例外）。
    这条测试的用处是：谁把它删了/移走了，测试会立刻红——而不是等到某天
    有人发现"对照校验一直在静默跳过"。
    """
    path = cr.bundled_template_path()
    assert path.is_file(), f"仓库里没有公司模版：{path}"
    assert path.stat().st_size > 100_000, "模版文件小得不像话，可能被截断了"


def test_bundled_template_has_the_three_expected_sheets():
    """参照物的形状就是契约：三张表的表名与顺序。改模版的人会先碰到这条。"""
    import openpyxl

    workbook = openpyxl.load_workbook(cr.bundled_template_path())
    assert workbook.sheetnames == [cr.SHEET_PLAN, cr.SHEET_LIST, cr.SHEET_SHOTS]


def test_compare_defaults_to_the_bundled_template(fixture, tmp_path):
    """`compare_with_template` 不给路径时**默认用仓库里那份**——这才是新人拉下来就能验的形态。"""
    model = build_model(fixture.results, fixture.cases)
    ours = write_report(model, tmp_path / "报告.xlsx")

    checks = compare_with_template(ours)          # 不传模版
    assert len(checks) >= 13
    assert any(check.name.startswith("「测试截图」表头值") for check in checks)
